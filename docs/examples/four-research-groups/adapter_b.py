# Copyright (c) 2026 Research Center for Energy Networks - ETH Zürich
# The CESDM project is part of the SWEET-CoSi consortium, funded by the
# Swiss Federal Office of Energy (SFOE) through its SWEET programme.

"""Group B: Excel workbook in a PyPSA-like style.

    Regions     region_id, region_name
    Generators  unit_id, region_id, p_nom, carrier, efficiency, marginal_cost
                MW; carrier = fuel name; efficiency is a fraction, blank if unknown
    Demand      region_id, demand_gwh                  GWh/year
    Storage     store_id, region_id, energy_gwh, inflow_gwh

Hydro, reservoir and run-of-river, is written here as a generator on a bus
of water: HydroGenerationUnit plus HydroReservoir.
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl


def _toolbox_root() -> Path:
    start = Path(__file__).resolve().parent
    for folder in (start, *start.parents):
        if (folder / "cesdm" / "schemas" / "base").is_dir():
            return folder
    raise FileNotFoundError("cesdm schema directory not found")


_ROOT = _toolbox_root()
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from cesdm_toolbox import build_model_from_yaml  # noqa: E402
from ear import get_attr_value  # noqa: E402

SOURCE = "group-b"


def network():
    """An empty CESDM model with the libraries Group B's workbook refers to."""
    model = build_model_from_yaml(str(_ROOT / "cesdm/schemas/base"))
    model.import_library(str(_ROOT / "cesdm/library/default_library"))
    model.import_library(str(_ROOT / "cesdm/library/regions_library"))
    return model


def network_from_yaml(path):
    model = network()
    model.import_yaml(str(path))
    return model


def cell(x: float):
    x = round(float(x), 9)
    return int(x) if float(x).is_integer() else x

REGIONS = {"CH": ("region.country.CH", "Switzerland")}

# Group B names fuels. The name becomes a category on the plant. Efficiency and
# marginal cost are columns of the workbook, not values of a library technology.
CARRIER = {
    "gas": "ccgt",
    "nuclear": "nuclear",
    "wind": "wind",
    "solar": "solar",
    "hydro": "reservoir_hydro",
    "run_of_river": "run_of_river",
}
TECH_TO_CARRIER = {tech: c for c, tech in CARRIER.items()}
REGION_TO_CODE = {cesdm: code for code, (cesdm, _) in REGIONS.items()}

_COMPONENT = {
    "wind": "WindGenerationUnit",
    "solar": "SolarGenerationUnit",
}


def _rows(ws) -> list[dict]:
    header, *rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(header, r)) for r in rows]


def to_cesdm(xlsx: Path, model) -> None:
    wb = openpyxl.load_workbook(xlsx, read_only=True)

    for row in _rows(wb["Demand"]):
        region = REGIONS[row["region_id"]][0]
        eid = f"demand.{row['region_id']}"
        model.add_entity("DemandUnit", eid)
        model.add_relation(eid, "belongsToGeographicalRegion", region)
        model.add_attribute(eid, "annual_energy_demand", float(row["demand_gwh"]) * 1000,
                            unit="MWh/year", provenance_ref=SOURCE)  # GWh -> MWh

    # Like Group A, the store belongs to the hydro unit of the same region by
    # convention: energy > 0 for the reservoir, energy = 0 for run-of-river.
    stores = {}
    for row in _rows(wb["Storage"]):
        kind = "hydro" if float(row["energy_gwh"]) > 0 else "run_of_river"
        stores[(row["region_id"], kind)] = row

    for row in _rows(wb["Generators"]):
        region = REGIONS[row["region_id"]][0]
        carrier = row["carrier"]
        eid = f"gen.{row['unit_id']}"
        if carrier in ("hydro", "run_of_river"):
            model.add_entity("HydroGenerationUnit", eid)
        else:
            model.add_entity(_COMPONENT.get(carrier, "GenerationUnit"), eid)
        model.add_relation(eid, "belongsToGeographicalRegion", region)
        model.add_attribute(eid, "nominal_power_capacity", float(row["p_nom"]),
                            unit="MW", provenance_ref=SOURCE)
        model.add_attribute(eid, "generator_technology_type", CARRIER[carrier])
        if row["efficiency"] not in (None, ""):
            model.add_attribute(eid, "energy_conversion_efficiency", float(row["efficiency"]),
                                unit="fraction", provenance_ref=SOURCE)
        model.add_attribute(eid, "variable_operating_cost", float(row["marginal_cost"]),
                            unit="MU/MWh", provenance_ref=SOURCE)
        if carrier in ("hydro", "run_of_river"):
            s = stores[(row["region_id"], carrier)]
            basin = f"storage.{s['store_id']}"
            model.add_entity("HydroReservoir", basin)
            model.add_attribute(basin, "energy_storage_capacity", float(s["energy_gwh"]) * 1000,
                                unit="MWh", provenance_ref=SOURCE)
            model.add_attribute(basin, "annual_natural_inflow_energy", float(s["inflow_gwh"]) * 1000,
                                unit="MWh/year", provenance_ref=SOURCE)
            model.add_relation(eid, "drawsFromHydroReservoir", basin)


def from_cesdm(model, xlsx: Path) -> None:
    xlsx.parent.mkdir(parents=True, exist_ok=True)
    basins = model.get_entities("HydroReservoir")
    plants = []
    for component in ("GenerationUnit", "WindGenerationUnit", "SolarGenerationUnit", "HydroGenerationUnit"):
        for eid, gen in model.get_entities(component).items():
            row = {
                "id": eid.removeprefix("gen."),
                "region": gen.get_relation("belongsToGeographicalRegion"),
                "technology": get_attr_value(gen, "generator_technology_type"),
                "capacity_mw": get_attr_value(gen, "nominal_power_capacity"),
                "efficiency": get_attr_value(gen, "energy_conversion_efficiency"),
                "variable_cost": get_attr_value(gen, "variable_operating_cost"),
            }
            sid = gen.get_relation("drawsFromHydroReservoir")
            if sid:
                basin = basins[str(sid)]
                row["storage_id"] = str(sid).removeprefix("storage.")
                row["storage_mwh"] = get_attr_value(basin, "energy_storage_capacity")
                row["inflow_mwh_per_year"] = get_attr_value(basin, "annual_natural_inflow_energy")
            plants.append(row)
    demand = [
        {"region": unit.get_relation("belongsToGeographicalRegion"),
         "annual_mwh": get_attr_value(unit, "annual_energy_demand")}
        for unit in model.get_entities("DemandUnit").values()
    ]
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def sheet(name, header, rows):
        ws = wb.create_sheet(name)
        ws.append(header)
        for r in rows:
            ws.append(r)

    codes = sorted({REGION_TO_CODE[p["region"]] for p in plants})
    sheet("Regions", ["region_id", "region_name"], [[c, REGIONS[c][1]] for c in codes])
    sheet("Generators", ["unit_id", "region_id", "p_nom", "carrier", "efficiency", "marginal_cost"],
          [[p["id"], REGION_TO_CODE[p["region"]], cell(p["capacity_mw"]),
            TECH_TO_CARRIER[p["technology"]],
            "" if p["efficiency"] is None else cell(p["efficiency"]), cell(p["variable_cost"])]
           for p in plants])
    sheet("Demand", ["region_id", "demand_gwh"],
          [[REGION_TO_CODE[d["region"]], cell(d["annual_mwh"] / 1000)] for d in demand])
    sheet("Storage", ["store_id", "region_id", "energy_gwh", "inflow_gwh"],
          [[p["storage_id"], REGION_TO_CODE[p["region"]], cell(p["storage_mwh"] / 1000),
            cell(p["inflow_mwh_per_year"] / 1000)]
           for p in plants if "storage_id" in p])
    wb.save(xlsx)
